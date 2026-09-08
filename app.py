from flask import Flask, render_template, request, redirect, session
from database import get_connection
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from flask import send_file

app = Flask(__name__)
app.secret_key = "kunci_rahasia_inventaris"

@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":
        username = request.form["username"]
        password = request.form["password"]

        try:
            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute(
                """
                SELECT * FROM users
                WHERE username = %s AND password = %s
                """,
                (username, password)
            )

            user = cursor.fetchone()

            cursor.close()
            connection.close()

            if user:
                session["user_id"] = user[0]
                session["username"] = user[1]

                return redirect("/")

            return render_template(
                "login.html",
                error="Username atau Password salah"
            )

        except Exception as e:
            return f"Terjadi Kesalahan {e}"

    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect("/login")

@app.route("/")
def index():

    if "user_id" not in session:
        return redirect("/login")

    try:

        connection = get_connection()
        cursor = connection.cursor()

        # Total perangkat
        cursor.execute(
            "SELECT COUNT(*) FROM perangkat"
        )

        total_perangkat = cursor.fetchone()[0]


        # Perangkat kondisi baik
        cursor.execute(
            "SELECT COUNT(*) FROM perangkat WHERE kondisi = %s",
            ("Baik",)
        )

        perangkat_baik = cursor.fetchone()[0]


        # Perangkat dalam perbaikan
        cursor.execute(
            "SELECT COUNT(*) FROM perangkat WHERE kondisi = %s",
            ("Dalam Perbaikan",)
        )

        perangkat_perbaikan = cursor.fetchone()[0]


        # Perangkat rusak
        cursor.execute(
            "SELECT COUNT(*) FROM perangkat WHERE kondisi = %s",
            ("Rusak",)
        )

        perangkat_rusak = cursor.fetchone()[0]


        # Data perangkat
        cursor.execute("""
            SELECT *
            FROM perangkat
            ORDER BY kode_aset ASC
            LIMIT 5
        """)

        perangkat_terbaru = cursor.fetchall()


        cursor.close()
        connection.close()


        return render_template(
            "dashboard.html",
            total_perangkat=total_perangkat,
            perangkat_baik=perangkat_baik,
            perangkat_perbaikan=perangkat_perbaikan,
            perangkat_rusak=perangkat_rusak,
            perangkat_terbaru=perangkat_terbaru
        )


    except Exception as e:

        return f"Terjadi kesalahan: {e}"

@app.route("/perangkat")
def perangkat():

    if "user_id" not in session:
            return redirect("/login")
    
    try:
        connection = get_connection()
        cursor = connection.cursor()

        keyword = request.args.get("keyword", "")
        kondisi = request.args.get("kondisi", "")

        query = """
            SELECT *
            FROM perangkat
            WHERE (
                kode_aset LIKE %s
                OR nama_perangkat LIKE %s
                OR merek LIKE %s
                OR lokasi LIKE %s
            )
        """

        keyword_search = "%" + keyword + "%"

        values = [
            keyword_search,
            keyword_search,
            keyword_search,
            keyword_search
        ]

        if kondisi:
            query += " AND kondisi = %s"
            values.append(kondisi)

        query += " ORDER BY kode_aset ASC"

        cursor.execute(query, values)
        data = cursor.fetchall()

        cursor.close()
        connection.close()

        return render_template(
            "perangkat.html",
            perangkat_list=data,
            keyword=keyword,
            kondisi=kondisi
        )

    except Exception as e:
        return f"Terjadi kesalahan: {e}"

@app.route("/perangkat/tambah", methods=["GET", "POST"])
def tambah_perangkat():

    if "user_id" not in session:
            return redirect("/login")

    if request.method == "POST":

        kode_aset = request.form["kode_aset"]
        nama_perangkat = request.form["nama_perangkat"]
        kategori = request.form["kategori"]
        merek = request.form["merek"]
        lokasi = request.form["lokasi"]
        kondisi = request.form["kondisi"]
        tanggal_masuk = request.form["tanggal_masuk"]
        keterangan = request.form["keterangan"]

        try:
            connection = get_connection()
            cursor = connection.cursor()

            quary = """
                INSERT INTO perangkat
                (kode_aset, nama_perangkat, kategori, merek, lokasi, kondisi, tanggal_masuk, keterangan)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            """

            values = (
                kode_aset,
                nama_perangkat,
                kategori,
                merek,
                lokasi,
                kondisi,
                tanggal_masuk,
                keterangan
            )

            cursor.execute(quary, values)
            connection.commit()

            cursor.close()
            connection.close()

            return redirect("/perangkat")

        except Exception as e:
            return f"Terjadi kesalahan: {e}"

    return render_template("tambah_perangkat.html")

@app.route("/perangkat/edit/<int:id>", methods=["GET", "POST"])
def edit_perangkat(id):

    try:
        connection = get_connection()
        cursor = connection.cursor()

        if request.method == "POST":

            kode_aset = request.form["kode_aset"]
            nama_perangkat = request.form["nama_perangkat"]
            kategori = request.form["kategori"]
            merek = request.form["merek"]
            lokasi = request.form["lokasi"]
            kondisi = request.form["kondisi"]
            tanggal_masuk = request.form["tanggal_masuk"]
            keterangan = request.form["keterangan"]

            query = """
                UPDATE perangkat
                SET kode_aset = %s,
                    nama_perangkat = %s,
                    kategori = %s,
                    merek = %s,
                    lokasi = %s,
                    kondisi = %s,
                    tanggal_masuk = %s,
                    keterangan = %s
                WHERE id = %s
            """

            values = (
                kode_aset,
                nama_perangkat,
                kategori,
                merek,
                lokasi,
                kondisi,
                tanggal_masuk,
                keterangan,
                id
            )

            cursor.execute(query, values)
            connection.commit()

            cursor.close()
            connection.close()

            return redirect("/perangkat")

        # Mengambil data bedasrkan ID
        cursor.execute(
            "SELECT * FROM perangkat WHERE id = %s",
            (id,)
        )

        data = cursor.fetchone()

        cursor.close()
        connection.close()

        if data is None:
            return "Data perangkat tidak ditemukan"

        return render_template(
            "edit_perangkat.html",
            perangkat=data
        )

    except Exception as e:
        return f"Terjadi Kesalahan: {e}"


@app.route("/perangkat/hapus/<int:id>", methods=["POST"])
def hapus_perangkat(id):

    try:
        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute(
            "DELETE FROM  perangkat WHERE id = %s",
            (id,)
        )

        connection.commit()

        cursor.close()
        connection.close()

        return redirect("/perangkat")

    except Exception as e:
        return f"Terjadi kesalahan: {e}"

@app.route("/perangkat/cetak")
def cetak_perangkat():

    if "user_id" not in session:
        return redirect("/login")

    try:

        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute("""
            SELECT *
            FROM perangkat
            ORDER BY kode_aset ASC
        """)

        data = cursor.fetchall()

        cursor.close()
        connection.close()

        tanggal_cetak = datetime.now().strftime(
            "%d-%m-%Y %H:%M:%S"
        )

        return render_template(
            "cetak_perangkat.html",
            perangkat_list=data,
            tanggal_cetak=tanggal_cetak
        )

    except Exception as e:

        return f"Terjadi kesalahan: {e}"

@app.route("/perangkat/export")
def export_perangkat():
    try:
        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute("""
            SELECT *
            FROM perangkat
            ORDER BY kode_aset ASC
        """)

        data = cursor.fetchall()

        cursor.close()
        connection.close()

        # Membuat file Exel
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Inventaris IT"

        # Judul
        sheet["A1"] = "LAPORAN INVENTARIS IT"
        sheet["A1"].font = Font(bold=True, size=14)
        sheet["A1"].alignment = Alignment(horizontal="center")

        sheet.merge_cells("A1:I1")

        #Header tabel
        headers = [
            "No",
            "Kode Aset",
            "Nama Perangkat",
            "Kategori",
            "Merek",
            "Lokasi",
            "Kondisi",
            "Tanggal Masuk",
            "Keterangan"
        ]

        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=3, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        # Data Perangkat
        for row, perangkat in enumerate(data, start=4):
            sheet.cell(row=row, column=1).value = row - 3
            sheet.cell(row=row, column=2).value = perangkat[1]
            sheet.cell(row=row, column=3).value = perangkat[2]
            sheet.cell(row=row, column=4).value = perangkat[3]
            sheet.cell(row=row, column=5).value = perangkat[4]
            sheet.cell(row=row, column=6).value = perangkat[5]
            sheet.cell(row=row, column=7).value = perangkat[6]
            sheet.cell(row=row, column=8).value = perangkat[7]
            sheet.cell(row=row, column=9).value = perangkat[8]

        #Lebar kolom
        lebar_kolom = {
            "A": 8,
            "B": 15,
            "C": 25,
            "D": 18,
            "E": 18,
            "F": 20,
            "G": 20,
            "H": 18,
            "I": 30
        }

        for kolom, lebar in lebar_kolom.items():
            sheet.column_dimensions[kolom].width = lebar

        #Menyimpan workbook ke memory
        from io import BytesIO

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name="laporan_Inventaris.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Exception as e:
        return f"Terjadi kesalahan: {e}"


if __name__ == "__main__":
    app.run(debug=True)